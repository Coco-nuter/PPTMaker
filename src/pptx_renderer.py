"""将 DeckSpec 确定性渲染为可编辑的 PowerPoint 原生对象。"""

from collections.abc import Callable
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from models import (
    BaseSlideSpec,
    BulletsSlideSpec,
    ClosingSlideSpec,
    ComparisonSlideSpec,
    CoverSlideSpec,
    DeckSpec,
    MetricsSlideSpec,
    ProcessSlideSpec,
    SectionSlideSpec,
    SlideSpec,
    ThemeSpec,
    TimelineSlideSpec,
    TwoColumnSlideSpec,
)

SLIDE_WIDTH = Inches(13.333333)
SLIDE_HEIGHT = Inches(7.5)
HORIZONTAL_MARGIN = Inches(0.8)
VERTICAL_MARGIN = Inches(0.55)
CONTENT_WIDTH = SLIDE_WIDTH - (2 * HORIZONTAL_MARGIN)
CONTENT_BOTTOM = SLIDE_HEIGHT - VERTICAL_MARGIN

def _rgb(color: str) -> RGBColor:
    """将经过 DeckSpec 校验的 #RRGGBB 转成 PowerPoint 颜色。"""
    return RGBColor.from_string(color.removeprefix("#"))


def _blend(color: str, target: str, target_ratio: float) -> str:
    """按固定比例混合两种主题色，用于确定性浅色背景。"""
    source_rgb = tuple(int(color[index : index + 2], 16) for index in (1, 3, 5))
    target_rgb = tuple(int(target[index : index + 2], 16) for index in (1, 3, 5))
    mixed = tuple(
        round(source + (destination - source) * target_ratio)
        for source, destination in zip(source_rgb, target_rgb, strict=True)
    )
    return f"#{mixed[0]:02X}{mixed[1]:02X}{mixed[2]:02X}"


def _require_slide[TSlide: BaseSlideSpec](
    spec: SlideSpec,
    expected_type: type[TSlide],
) -> TSlide:
    """保护布局分发表，避免错误类型被静默渲染。"""
    if not isinstance(spec, expected_type):
        raise TypeError(f"layout {spec.layout!r} requires {expected_type.__name__}")
    return spec


def _set_background(slide: Slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _add_rectangle(
    slide: Slide,
    *,
    name: str,
    left: int,
    top: int,
    width: int,
    height: int,
    color: str,
    border_color: str | None = None,
    border_width: float = 1.0,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    if border_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(border_color)
        shape.line.width = Pt(border_width)


def _add_rounded_card(
    slide: Slide,
    *,
    name: str,
    left: int,
    top: int,
    width: int,
    height: int,
    color: str,
    border_color: str,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.color.rgb = _rgb(border_color)
    shape.line.width = Pt(1.0)


def _add_textbox(
    slide: Slide,
    *,
    name: str,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    font_name: str,
    font_size: int,
    color: str,
    bold: bool = False,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
    vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.vertical_anchor = vertical_anchor

    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.line_spacing = 1.0
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _add_circle_node(
    slide: Slide,
    *,
    name: str,
    left: int,
    top: int,
    diameter: int,
    color: str,
    border_color: str | None = None,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    if border_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(border_color)
        shape.line.width = Pt(1.0)


def _add_connector(
    slide: Slide,
    *,
    name: str,
    begin_x: int,
    begin_y: int,
    end_x: int,
    end_y: int,
    color: str,
    width: float = 2.0,
) -> None:
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        begin_x,
        begin_y,
        end_x,
        end_y,
    )
    connector.name = name
    connector.line.color.rgb = _rgb(color)
    connector.line.width = Pt(width)


def _add_arrow(
    slide: Slide,
    *,
    name: str,
    left: int,
    top: int,
    width: int,
    height: int,
    color: str,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, width, height)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()


def _add_number_badge(
    slide: Slide,
    *,
    name: str,
    number: str,
    left: int,
    top: int,
    diameter: int,
    fill_color: str,
    text_color: str,
    font_name: str,
    font_size: int = 16,
) -> None:
    _add_circle_node(
        slide,
        name=f"{name} Circle",
        left=left,
        top=top,
        diameter=diameter,
        color=fill_color,
    )
    _add_textbox(
        slide,
        name=f"{name} Text",
        text=number,
        left=left,
        top=top,
        width=diameter,
        height=diameter,
        font_name=font_name,
        font_size=font_size,
        color=text_color,
        bold=True,
        alignment=PP_ALIGN.CENTER,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )


def _add_page_title(slide: Slide, title: str, theme: ThemeSpec, *, name: str) -> None:
    """绘制所有内容页共享的标题和装饰色块。"""
    _add_textbox(
        slide,
        name=f"{name} Title",
        text=title,
        left=HORIZONTAL_MARGIN,
        top=VERTICAL_MARGIN,
        width=Inches(10.9),
        height=Inches(0.72),
        font_name=theme.heading_font,
        font_size=35,
        color=theme.primary_color,
        bold=True,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    _add_rectangle(
        slide,
        name=f"{name} Title Accent",
        left=HORIZONTAL_MARGIN,
        top=Inches(1.33),
        width=Inches(0.95),
        height=Inches(0.06),
        color=theme.secondary_color,
    )
    _add_rectangle(
        slide,
        name=f"{name} Corner Accent",
        left=Inches(11.88),
        top=VERTICAL_MARGIN,
        width=Inches(0.65),
        height=Inches(0.16),
        color=theme.secondary_color,
    )


def _render_cover(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, CoverSlideSpec)
    _set_background(slide, theme.background_color)
    _add_rectangle(
        slide,
        name="Cover Primary Block",
        left=HORIZONTAL_MARGIN,
        top=VERTICAL_MARGIN,
        width=Inches(0.55),
        height=Inches(6.4),
        color=theme.primary_color,
    )
    _add_rectangle(
        slide,
        name="Cover Bottom Accent",
        left=Inches(1.75),
        top=Inches(5.72),
        width=Inches(3.7),
        height=Inches(0.09),
        color=theme.secondary_color,
    )
    _add_rectangle(
        slide,
        name="Cover Decorative Block",
        left=Inches(10.05),
        top=Inches(4.95),
        width=Inches(2.48),
        height=Inches(1.55),
        color=_blend(theme.secondary_color, theme.background_color, 0.72),
    )
    _add_circle_node(
        slide,
        name="Cover Decorative Circle Large",
        left=Inches(10.25),
        top=Inches(1.15),
        diameter=Inches(1.32),
        color=_blend(theme.primary_color, theme.background_color, 0.78),
    )
    _add_circle_node(
        slide,
        name="Cover Decorative Circle Small",
        left=Inches(11.52),
        top=Inches(2.25),
        diameter=Inches(0.7),
        color=theme.secondary_color,
    )
    _add_textbox(
        slide,
        name="Cover Title",
        text=spec.title,
        left=Inches(1.75),
        top=Inches(1.55),
        width=Inches(8.15),
        height=Inches(1.55),
        font_name=theme.heading_font,
        font_size=52,
        color=theme.primary_color,
        bold=True,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    if spec.subtitle:
        _add_textbox(
            slide,
            name="Cover Subtitle",
            text=spec.subtitle,
            left=Inches(1.78),
            top=Inches(3.45),
            width=Inches(7.85),
            height=Inches(1.0),
            font_name=theme.body_font,
            font_size=24,
            color=theme.text_color,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )


def _render_section(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, SectionSlideSpec)
    _set_background(slide, _blend(theme.primary_color, theme.background_color, 0.93))
    _add_rectangle(
        slide,
        name="Section Divider",
        left=Inches(3.52),
        top=Inches(1.25),
        width=Inches(0.08),
        height=Inches(4.9),
        color=theme.secondary_color,
    )
    _add_rectangle(
        slide,
        name="Section Bottom Accent",
        left=Inches(4.12),
        top=Inches(5.45),
        width=Inches(3.65),
        height=Inches(0.1),
        color=theme.secondary_color,
    )
    _add_circle_node(
        slide,
        name="Section Decorative Node",
        left=Inches(11.45),
        top=Inches(0.85),
        diameter=Inches(0.75),
        color=_blend(theme.secondary_color, theme.background_color, 0.35),
    )
    _add_textbox(
        slide,
        name="Section Number",
        text=spec.section_number or "•",
        left=Inches(1.0),
        top=Inches(1.65),
        width=Inches(2.1),
        height=Inches(1.55),
        font_name=theme.heading_font,
        font_size=72,
        color=theme.secondary_color,
        bold=True,
        alignment=PP_ALIGN.CENTER,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    _add_textbox(
        slide,
        name="Section Title",
        text=spec.title,
        left=Inches(4.12),
        top=Inches(1.9),
        width=Inches(7.65),
        height=Inches(1.15),
        font_name=theme.heading_font,
        font_size=44,
        color=theme.primary_color,
        bold=True,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    _add_textbox(
        slide,
        name="Section Subtitle",
        text=spec.subtitle,
        left=Inches(4.15),
        top=Inches(3.35),
        width=Inches(7.35),
        height=Inches(0.9),
        font_name=theme.body_font,
        font_size=24,
        color=theme.text_color,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )


def _render_bullets(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, BulletsSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Bullets")

    card_top = Inches(1.72)
    if spec.subtitle:
        _add_textbox(
            slide,
            name="Bullets Subtitle",
            text=spec.subtitle,
            left=HORIZONTAL_MARGIN,
            top=Inches(1.48),
            width=CONTENT_WIDTH,
            height=Inches(0.35),
            font_name=theme.body_font,
            font_size=16,
            color=theme.secondary_color,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        card_top = Inches(1.92)

    count = len(spec.bullets)
    gap = Inches(0.12)
    available_height = CONTENT_BOTTOM - card_top
    card_height = min(
        Inches(1.2),
        (available_height - gap * (count - 1)) // count,
    )
    light_card = _blend(theme.primary_color, theme.background_color, 0.92)
    badge_size = Inches(0.42)

    for index, bullet in enumerate(spec.bullets, start=1):
        top = card_top + (index - 1) * (card_height + gap)
        _add_rounded_card(
            slide,
            name=f"Bullet Card {index}",
            left=HORIZONTAL_MARGIN,
            top=top,
            width=CONTENT_WIDTH,
            height=card_height,
            color=light_card,
            border_color=_blend(theme.primary_color, theme.background_color, 0.72),
        )
        badge_top = top + (card_height - badge_size) // 2
        _add_number_badge(
            slide,
            name=f"Bullet Number {index}",
            number=f"{index:02d}",
            left=Inches(1.05),
            top=badge_top,
            diameter=badge_size,
            fill_color=theme.secondary_color,
            text_color=theme.background_color,
            font_name=theme.heading_font,
            font_size=16,
        )
        _add_textbox(
            slide,
            name=f"Bullet Text {index}",
            text=bullet,
            left=Inches(1.72),
            top=top + Inches(0.08),
            width=Inches(10.35),
            height=card_height - Inches(0.16),
            font_name=theme.body_font,
            font_size=18 if count <= 4 else 16,
            color=theme.text_color,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )


def _render_two_column(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, TwoColumnSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Two Column")

    panel_top = Inches(1.62)
    panel_height = Inches(5.15)
    gap = Inches(0.34)
    panel_width = (CONTENT_WIDTH - gap) // 2
    left_positions = (
        HORIZONTAL_MARGIN,
        HORIZONTAL_MARGIN + panel_width + gap,
    )
    columns = (
        (spec.left_title, spec.left_items, theme.primary_color, "Left"),
        (spec.right_title, spec.right_items, theme.secondary_color, "Right"),
    )

    for left, (title, items, accent_color, side) in zip(left_positions, columns, strict=True):
        light_color = _blend(accent_color, theme.background_color, 0.91)
        _add_rounded_card(
            slide,
            name=f"Two Column {side} Panel",
            left=left,
            top=panel_top,
            width=panel_width,
            height=panel_height,
            color=light_color,
            border_color=_blend(accent_color, theme.background_color, 0.58),
        )
        _add_rectangle(
            slide,
            name=f"Two Column {side} Header",
            left=left,
            top=panel_top,
            width=panel_width,
            height=Inches(0.78),
            color=accent_color,
        )
        _add_textbox(
            slide,
            name=f"Two Column {side} Title",
            text=title,
            left=left + Inches(0.28),
            top=panel_top + Inches(0.1),
            width=panel_width - Inches(0.56),
            height=Inches(0.56),
            font_name=theme.heading_font,
            font_size=24,
            color=theme.background_color,
            bold=True,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        item_top = panel_top + Inches(1.08)
        item_height = (panel_height - Inches(1.35)) // len(items)
        for index, item in enumerate(items, start=1):
            top = item_top + (index - 1) * item_height
            _add_circle_node(
                slide,
                name=f"Two Column {side} Bullet {index}",
                left=left + Inches(0.32),
                top=top + Inches(0.16),
                diameter=Inches(0.14),
                color=accent_color,
            )
            _add_textbox(
                slide,
                name=f"Two Column {side} Item {index}",
                text=item,
                left=left + Inches(0.62),
                top=top,
                width=panel_width - Inches(0.92),
                height=item_height,
                font_name=theme.body_font,
                font_size=16,
                color=theme.text_color,
                vertical_anchor=MSO_ANCHOR.MIDDLE,
            )


def _render_metrics(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, MetricsSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Metrics")

    count = len(spec.metrics)
    gap = Inches(0.22)
    card_width = (CONTENT_WIDTH - gap * (count - 1)) // count
    card_top = Inches(1.82)
    card_height = Inches(4.72)

    for index, metric in enumerate(spec.metrics, start=1):
        left = HORIZONTAL_MARGIN + (index - 1) * (card_width + gap)
        accent_color = theme.primary_color if index % 2 else theme.secondary_color
        _add_rounded_card(
            slide,
            name=f"Metric Card {index}",
            left=left,
            top=card_top,
            width=card_width,
            height=card_height,
            color=_blend(accent_color, theme.background_color, 0.91),
            border_color=_blend(accent_color, theme.background_color, 0.56),
        )
        _add_rectangle(
            slide,
            name=f"Metric Accent {index}",
            left=left,
            top=card_top,
            width=card_width,
            height=Inches(0.12),
            color=accent_color,
        )
        _add_textbox(
            slide,
            name=f"Metric Value {index}",
            text=metric.value,
            left=left + Inches(0.22),
            top=Inches(2.25),
            width=card_width - Inches(0.44),
            height=Inches(0.9),
            font_name=theme.heading_font,
            font_size=34,
            color=accent_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_textbox(
            slide,
            name=f"Metric Label {index}",
            text=metric.label,
            left=left + Inches(0.22),
            top=Inches(3.4),
            width=card_width - Inches(0.44),
            height=Inches(0.6),
            font_name=theme.heading_font,
            font_size=20,
            color=theme.text_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_textbox(
            slide,
            name=f"Metric Description {index}",
            text=metric.description,
            left=left + Inches(0.28),
            top=Inches(4.28),
            width=card_width - Inches(0.56),
            height=Inches(1.35),
            font_name=theme.body_font,
            font_size=16,
            color=theme.text_color,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )


def _render_timeline(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, TimelineSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Timeline")

    count = len(spec.items)
    slot_width = CONTENT_WIDTH // count
    node_size = Inches(0.48)
    line_y = Inches(3.0)
    first_center = HORIZONTAL_MARGIN + slot_width // 2
    last_center = HORIZONTAL_MARGIN + (count - 1) * slot_width + slot_width // 2
    _add_connector(
        slide,
        name="Timeline Connector",
        begin_x=first_center,
        begin_y=line_y,
        end_x=last_center,
        end_y=line_y,
        color=_blend(theme.primary_color, theme.background_color, 0.35),
        width=3.0,
    )

    for index, item in enumerate(spec.items, start=1):
        slot_left = HORIZONTAL_MARGIN + (index - 1) * slot_width
        center = slot_left + slot_width // 2
        accent_color = theme.primary_color if index % 2 else theme.secondary_color
        _add_textbox(
            slide,
            name=f"Timeline Label {index}",
            text=item.label,
            left=slot_left + Inches(0.08),
            top=Inches(1.9),
            width=slot_width - Inches(0.16),
            height=Inches(0.5),
            font_name=theme.body_font,
            font_size=16,
            color=accent_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_number_badge(
            slide,
            name=f"Timeline Node {index}",
            number=str(index),
            left=center - node_size // 2,
            top=line_y - node_size // 2,
            diameter=node_size,
            fill_color=accent_color,
            text_color=theme.background_color,
            font_name=theme.heading_font,
            font_size=16,
        )
        _add_textbox(
            slide,
            name=f"Timeline Item Title {index}",
            text=item.title,
            left=slot_left + Inches(0.08),
            top=Inches(3.52),
            width=slot_width - Inches(0.16),
            height=Inches(0.65),
            font_name=theme.heading_font,
            font_size=20,
            color=theme.text_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_textbox(
            slide,
            name=f"Timeline Description {index}",
            text=item.description,
            left=slot_left + Inches(0.1),
            top=Inches(4.28),
            width=slot_width - Inches(0.2),
            height=Inches(1.45),
            font_name=theme.body_font,
            font_size=16,
            color=theme.text_color,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.TOP,
        )


def _render_process(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, ProcessSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Process")

    count = len(spec.steps)
    gap = Inches(0.28)
    card_width = (CONTENT_WIDTH - gap * (count - 1)) // count
    card_top = Inches(1.9)
    card_height = Inches(4.65)

    for index in range(1, count):
        arrow_left = HORIZONTAL_MARGIN + index * card_width + (index - 1) * gap
        _add_arrow(
            slide,
            name=f"Process Arrow {index}",
            left=arrow_left + Inches(0.055),
            top=Inches(3.72),
            width=gap - Inches(0.11),
            height=Inches(0.36),
            color=_blend(theme.secondary_color, theme.background_color, 0.18),
        )

    for index, step in enumerate(spec.steps, start=1):
        left = HORIZONTAL_MARGIN + (index - 1) * (card_width + gap)
        accent_color = theme.primary_color if index % 2 else theme.secondary_color
        _add_rounded_card(
            slide,
            name=f"Process Card {index}",
            left=left,
            top=card_top,
            width=card_width,
            height=card_height,
            color=_blend(accent_color, theme.background_color, 0.92),
            border_color=_blend(accent_color, theme.background_color, 0.58),
        )
        badge_size = Inches(0.56)
        _add_number_badge(
            slide,
            name=f"Process Number {index}",
            number=f"{index:02d}",
            left=left + (card_width - badge_size) // 2,
            top=Inches(2.2),
            diameter=badge_size,
            fill_color=accent_color,
            text_color=theme.background_color,
            font_name=theme.heading_font,
            font_size=16,
        )
        _add_textbox(
            slide,
            name=f"Process Step Title {index}",
            text=step.title,
            left=left + Inches(0.14),
            top=Inches(3.08),
            width=card_width - Inches(0.28),
            height=Inches(0.72),
            font_name=theme.heading_font,
            font_size=20,
            color=theme.text_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_textbox(
            slide,
            name=f"Process Description {index}",
            text=step.description,
            left=left + Inches(0.16),
            top=Inches(4.0),
            width=card_width - Inches(0.32),
            height=Inches(1.55),
            font_name=theme.body_font,
            font_size=16,
            color=theme.text_color,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.TOP,
        )


def _render_comparison(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, ComparisonSlideSpec)
    _set_background(slide, theme.background_color)
    _add_page_title(slide, spec.title, theme, name="Comparison")

    label_width = Inches(1.75)
    column_gap = Inches(0.14)
    value_width = (CONTENT_WIDTH - label_width - column_gap * 2) // 2
    label_left = HORIZONTAL_MARGIN
    left_value_left = label_left + label_width + column_gap
    right_value_left = left_value_left + value_width + column_gap
    header_top = Inches(1.65)
    header_height = Inches(0.72)

    headers = (
        (label_left, label_width, "对比维度", theme.text_color, theme.background_color),
        (
            left_value_left,
            value_width,
            spec.left_title,
            theme.background_color,
            theme.primary_color,
        ),
        (
            right_value_left,
            value_width,
            spec.right_title,
            theme.background_color,
            theme.secondary_color,
        ),
    )
    for index, (left, width, text, text_color, fill_color) in enumerate(headers, start=1):
        _add_rounded_card(
            slide,
            name=f"Comparison Header {index}",
            left=left,
            top=header_top,
            width=width,
            height=header_height,
            color=fill_color,
            border_color=_blend(fill_color, theme.background_color, 0.45),
        )
        _add_textbox(
            slide,
            name=f"Comparison Header Title {index}",
            text=text,
            left=left + Inches(0.12),
            top=header_top + Inches(0.08),
            width=width - Inches(0.24),
            height=header_height - Inches(0.16),
            font_name=theme.heading_font,
            font_size=18 if index == 1 else 22,
            color=text_color,
            bold=True,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )

    row_top = Inches(2.58)
    row_gap = Inches(0.08)
    row_height = min(
        Inches(0.76),
        (CONTENT_BOTTOM - row_top - row_gap * (len(spec.rows) - 1)) // len(spec.rows),
    )
    left_fill = _blend(theme.primary_color, theme.background_color, 0.92)
    right_fill = _blend(theme.secondary_color, theme.background_color, 0.9)
    label_fill = _blend(theme.text_color, theme.background_color, 0.94)

    for index, row in enumerate(spec.rows, start=1):
        top = row_top + (index - 1) * (row_height + row_gap)
        cells = (
            (label_left, label_width, row.label, label_fill, True),
            (left_value_left, value_width, row.left_value, left_fill, False),
            (right_value_left, value_width, row.right_value, right_fill, False),
        )
        for cell_index, (left, width, text, fill_color, bold) in enumerate(cells, start=1):
            _add_rectangle(
                slide,
                name=f"Comparison Row {index} Cell {cell_index}",
                left=left,
                top=top,
                width=width,
                height=row_height,
                color=fill_color,
                border_color=_blend(theme.text_color, theme.background_color, 0.82),
                border_width=0.75,
            )
            _add_textbox(
                slide,
                name=f"Comparison Row {index} Text {cell_index}",
                text=text,
                left=left + Inches(0.12),
                top=top + Inches(0.06),
                width=width - Inches(0.24),
                height=row_height - Inches(0.12),
                font_name=theme.heading_font if bold else theme.body_font,
                font_size=16,
                color=theme.text_color,
                bold=bold,
                alignment=PP_ALIGN.CENTER,
                vertical_anchor=MSO_ANCHOR.MIDDLE,
            )


def _render_closing(slide: Slide, raw_spec: SlideSpec, theme: ThemeSpec) -> None:
    spec = _require_slide(raw_spec, ClosingSlideSpec)
    _set_background(slide, theme.primary_color)
    _add_rectangle(
        slide,
        name="Closing Top Accent",
        left=HORIZONTAL_MARGIN,
        top=VERTICAL_MARGIN,
        width=Inches(1.85),
        height=Inches(0.13),
        color=theme.secondary_color,
    )
    _add_rectangle(
        slide,
        name="Closing Bottom Accent",
        left=Inches(10.1),
        top=Inches(6.72),
        width=Inches(2.43),
        height=Inches(0.2),
        color=theme.secondary_color,
    )
    _add_circle_node(
        slide,
        name="Closing Decorative Circle Large",
        left=Inches(1.0),
        top=Inches(4.95),
        diameter=Inches(1.22),
        color=_blend(theme.secondary_color, theme.primary_color, 0.35),
    )
    _add_circle_node(
        slide,
        name="Closing Decorative Circle Small",
        left=Inches(2.1),
        top=Inches(5.95),
        diameter=Inches(0.55),
        color=theme.secondary_color,
    )
    _add_textbox(
        slide,
        name="Closing Title",
        text=spec.title,
        left=Inches(1.2),
        top=Inches(1.7),
        width=Inches(10.9),
        height=Inches(1.45),
        font_name=theme.heading_font,
        font_size=50,
        color=theme.background_color,
        bold=True,
        alignment=PP_ALIGN.CENTER,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    _add_rectangle(
        slide,
        name="Closing Center Accent",
        left=Inches(5.55),
        top=Inches(3.48),
        width=Inches(2.23),
        height=Inches(0.08),
        color=theme.secondary_color,
    )
    if spec.subtitle:
        _add_textbox(
            slide,
            name="Closing Subtitle",
            text=spec.subtitle,
            left=Inches(1.35),
            top=Inches(3.9),
            width=Inches(10.65),
            height=Inches(0.9),
            font_name=theme.body_font,
            font_size=24,
            color=theme.background_color,
            alignment=PP_ALIGN.CENTER,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )


LayoutRenderer = Callable[[Slide, SlideSpec, ThemeSpec], None]
LAYOUT_RENDERERS: dict[str, LayoutRenderer] = {
    "cover": _render_cover,
    "section": _render_section,
    "bullets": _render_bullets,
    "two_column": _render_two_column,
    "metrics": _render_metrics,
    "timeline": _render_timeline,
    "process": _render_process,
    "comparison": _render_comparison,
    "closing": _render_closing,
}


def render_deck(deck: DeckSpec, output_path: str | Path) -> Path:
    """渲染 DeckSpec，保存后重新打开并核对页数。"""
    destination = Path(output_path)
    if destination.suffix.lower() != ".pptx":
        raise ValueError("output path must use the .pptx extension")
    destination.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    presentation.core_properties.title = deck.title
    presentation.core_properties.subject = deck.purpose

    blank_layout = presentation.slide_layouts[6]
    for slide_spec in deck.slides:
        slide = presentation.slides.add_slide(blank_layout)
        LAYOUT_RENDERERS[slide_spec.layout](slide, slide_spec, deck.theme)

    presentation.save(destination)

    reopened = Presentation(destination)
    if len(reopened.slides) != len(deck.slides):
        raise RuntimeError("saved PPTX slide count does not match DeckSpec")

    return destination
