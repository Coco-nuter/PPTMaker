"""DeckSpec 到九种可编辑 PPTX 布局的渲染测试。"""

import json
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from models import (
    BulletsSlideSpec,
    ClosingSlideSpec,
    ComparisonSlideSpec,
    CoverSlideSpec,
    DeckSpec,
    MetricsSlideSpec,
    ProcessSlideSpec,
    SectionSlideSpec,
    SlideSpec,
    TimelineSlideSpec,
    TwoColumnSlideSpec,
)
from pptx_renderer import (
    HORIZONTAL_MARGIN,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    VERTICAL_MARGIN,
    render_deck,
)

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "sample_deck.json"
EXPECTED_LAYOUTS = [
    "cover",
    "section",
    "bullets",
    "two_column",
    "metrics",
    "timeline",
    "process",
    "comparison",
    "closing",
]


def load_sample_deck() -> DeckSpec:
    return DeckSpec.model_validate_json(FIXTURE_PATH.read_text(encoding="utf-8"))


def paragraph_texts(presentation: Presentation) -> list[str]:
    return [
        paragraph.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text
    ]


def expected_slide_texts(slide: SlideSpec) -> list[str]:
    texts = [slide.title]
    if isinstance(slide, (CoverSlideSpec, ClosingSlideSpec)):
        return texts + ([slide.subtitle] if slide.subtitle else [])
    if isinstance(slide, SectionSlideSpec):
        return texts + ([slide.section_number] if slide.section_number else []) + [slide.subtitle]
    if isinstance(slide, BulletsSlideSpec):
        return texts + ([slide.subtitle] if slide.subtitle else []) + slide.bullets
    if isinstance(slide, TwoColumnSlideSpec):
        return texts + [slide.left_title, *slide.left_items, slide.right_title, *slide.right_items]
    if isinstance(slide, MetricsSlideSpec):
        return texts + [
            value
            for metric in slide.metrics
            for value in (metric.value, metric.label, metric.description)
        ]
    if isinstance(slide, TimelineSlideSpec):
        return texts + [
            value
            for item in slide.items
            for value in (item.label, item.title, item.description)
        ]
    if isinstance(slide, ProcessSlideSpec):
        return texts + [
            value for step in slide.steps for value in (step.title, step.description)
        ]
    if isinstance(slide, ComparisonSlideSpec):
        return texts + [slide.left_title, slide.right_title] + [
            value
            for row in slide.rows
            for value in (row.label, row.left_value, row.right_value)
        ]
    raise AssertionError(f"unexpected slide type: {type(slide).__name__}")


def structural_summary(presentation: Presentation) -> list[tuple[object, ...]]:
    return [
        (
            shape.name,
            shape.shape_type,
            shape.left,
            shape.top,
            shape.width,
            shape.height,
            tuple(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if shape.has_text_frame
            else (),
        )
        for slide in presentation.slides
        for shape in slide.shapes
    ]


def test_nine_layout_deck_renders_and_reopens_without_mutation(tmp_path: Path) -> None:
    deck = load_sample_deck()
    original_json = deck.model_dump_json()
    output_path = tmp_path / "sample_deck.pptx"

    result = render_deck(deck, output_path)
    reopened = Presentation(result)

    assert result == output_path
    assert result.is_file()
    assert len(reopened.slides) == 9
    assert [slide.layout for slide in deck.slides] == EXPECTED_LAYOUTS
    assert reopened.slide_width == SLIDE_WIDTH
    assert reopened.slide_height == SLIDE_HEIGHT
    assert deck.model_dump_json() == original_json
    assert all(
        shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        for slide in reopened.slides
        for shape in slide.shapes
    )
    with ZipFile(result) as archive:
        assert not any(name.startswith("ppt/media/") for name in archive.namelist())


def test_rendered_text_matches_all_layout_content(tmp_path: Path) -> None:
    deck = load_sample_deck()
    presentation = Presentation(render_deck(deck, tmp_path / "content.pptx"))

    for slide_spec, rendered_slide in zip(deck.slides, presentation.slides, strict=True):
        rendered_text = [
            paragraph.text
            for shape in rendered_slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text
        ]
        for expected_text in expected_slide_texts(slide_spec):
            assert expected_text in rendered_text


def test_layouts_use_expected_native_shape_structures(tmp_path: Path) -> None:
    deck = load_sample_deck()
    presentation = Presentation(render_deck(deck, tmp_path / "structures.pptx"))
    names = [[shape.name for shape in slide.shapes] for slide in presentation.slides]

    assert "Cover Decorative Circle Large" in names[0]
    assert "Section Number" in names[1]
    assert sum(name.startswith("Bullet Card ") for name in names[2]) == 3
    assert "Two Column Left Panel" in names[3]
    assert "Two Column Right Panel" in names[3]
    assert sum(name.startswith("Metric Card ") for name in names[4]) == 3
    assert sum(name.endswith(" Circle") and name.startswith("Timeline Node") for name in names[5]) == 3
    assert sum(name.startswith("Process Card ") for name in names[6]) == 3
    assert sum(name.startswith("Comparison Row ") and " Cell " in name for name in names[7]) == 9
    assert "Closing Center Accent" in names[8]

    assert names[5].index("Timeline Connector") < names[5].index("Timeline Node 1 Circle")
    assert names[6].index("Process Arrow 1") < names[6].index("Process Card 1")


def test_shapes_stay_inside_safe_area_and_text_uses_theme(tmp_path: Path) -> None:
    deck = load_sample_deck()
    presentation = Presentation(render_deck(deck, tmp_path / "theme.pptx"))

    for slide_spec, slide in zip(deck.slides, presentation.slides, strict=True):
        title_found = False
        for shape in slide.shapes:
            assert shape.left >= HORIZONTAL_MARGIN
            assert shape.left + shape.width <= SLIDE_WIDTH - HORIZONTAL_MARGIN
            assert shape.top >= VERTICAL_MARGIN
            assert shape.top + shape.height <= SLIDE_HEIGHT - VERTICAL_MARGIN

            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font.name in {deck.theme.heading_font, deck.theme.body_font}
                    assert run.font.size is not None
                    assert run.font.size >= Pt(16)
                    if run.text == slide_spec.title:
                        title_found = True
                        assert run.font.name == deck.theme.heading_font
                        assert run.font.size >= Pt(35)
        assert title_found

    assert presentation.slides[0].background.fill.fore_color.rgb == RGBColor.from_string(
        deck.theme.background_color.removeprefix("#")
    )
    assert presentation.slides[8].background.fill.fore_color.rgb == RGBColor.from_string(
        deck.theme.primary_color.removeprefix("#")
    )


def test_metric_typography_has_clear_hierarchy(tmp_path: Path) -> None:
    presentation = Presentation(render_deck(load_sample_deck(), tmp_path / "metrics.pptx"))
    metric_slide = presentation.slides[4]

    def font_size(shape_name: str) -> int:
        shape = next(shape for shape in metric_slide.shapes if shape.name == shape_name)
        return shape.text_frame.paragraphs[0].runs[0].font.size

    assert font_size("Metric Value 1") > font_size("Metric Label 1")
    assert font_size("Metric Label 1") > font_size("Metric Description 1")


def test_layout_capacity_limits_render_inside_canvas(tmp_path: Path) -> None:
    data = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
    slides = data["slides"]
    slides[2]["bullets"] = [f"要点 {index}" for index in range(1, 7)]
    slides[3]["left_items"] = [f"左侧内容 {index}" for index in range(1, 7)]
    slides[3]["right_items"] = [f"右侧内容 {index}" for index in range(1, 7)]
    slides[4]["metrics"] = [
        {"value": f"{index}x", "label": f"指标 {index}", "description": "简短说明"}
        for index in range(1, 5)
    ]
    slides[5]["items"] = [
        {"label": f"阶段 {index}", "title": f"节点 {index}", "description": "简短说明"}
        for index in range(1, 7)
    ]
    slides[6]["steps"] = [
        {"title": f"步骤 {index}", "description": "简短说明"} for index in range(1, 7)
    ]
    slides[7]["rows"] = [
        {"label": f"维度 {index}", "left_value": "左侧", "right_value": "右侧"}
        for index in range(1, 7)
    ]
    deck = DeckSpec.model_validate(data)
    presentation = Presentation(render_deck(deck, tmp_path / "capacity.pptx"))

    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left >= HORIZONTAL_MARGIN
            assert shape.left + shape.width <= SLIDE_WIDTH - HORIZONTAL_MARGIN
            assert shape.top >= VERTICAL_MARGIN
            assert shape.top + shape.height <= SLIDE_HEIGHT - VERTICAL_MARGIN


def test_same_deck_produces_same_structure(tmp_path: Path) -> None:
    deck = load_sample_deck()
    first = Presentation(render_deck(deck, tmp_path / "first.pptx"))
    second = Presentation(render_deck(deck, tmp_path / "second.pptx"))

    assert structural_summary(first) == structural_summary(second)
