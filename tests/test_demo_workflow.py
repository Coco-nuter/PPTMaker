"""无 LLM Demo 业务编排测试。"""

from pathlib import Path
from unittest.mock import patch

import pytest
from pptx import Presentation

from demo_workflow import DemoGenerationError, generate_demo_project, load_deck_spec
from models import DeckSpec
from preview import PreviewError, PreviewResult

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "basic_deck.json"


def fake_preview(pptx_path: str | Path, output_root: str | Path) -> PreviewResult:
    source = Path(pptx_path).resolve()
    output_dir = Path(output_root).resolve() / "preview-test"
    png_dir = output_dir / "png"
    png_dir.mkdir(parents=True)
    png_paths = tuple(png_dir / f"slide_{index:03d}.png" for index in range(1, 4))
    for path in png_paths:
        path.write_bytes(b"mock png")
    return PreviewResult(
        source_pptx=source,
        output_dir=output_dir,
        png_paths=png_paths,
        page_count=3,
        width=1920,
        height=1080,
    )


def test_load_sample_deck_validates_fixture() -> None:
    deck = load_deck_spec(FIXTURE_PATH)

    assert isinstance(deck, DeckSpec)
    assert deck.title == "AI PPT Agent 项目介绍"
    assert len(deck.slides) == 3


def test_generation_creates_isolated_project_and_artifacts(tmp_path: Path) -> None:
    deck = load_deck_spec(FIXTURE_PATH)

    with patch("demo_workflow.pptx_to_pngs", side_effect=fake_preview):
        first = generate_demo_project(deck, tmp_path / "workspace")
        second = generate_demo_project(deck, tmp_path / "workspace")

    assert first.project_id != second.project_id
    assert first.project_dir.parent == (tmp_path / "workspace").resolve()
    assert first.deck_path.is_file()
    assert DeckSpec.model_validate_json(first.deck_path.read_text(encoding="utf-8")) == deck
    assert first.pptx_path.is_file()
    assert len(Presentation(first.pptx_path).slides) == 3
    assert len(first.preview_paths) == 3
    assert all(path.is_file() and first.project_dir in path.parents for path in first.preview_paths)


def test_failed_runs_never_reuse_project_directory(tmp_path: Path) -> None:
    deck = load_deck_spec(FIXTURE_PATH)

    with patch(
        "demo_workflow.pptx_to_pngs",
        side_effect=PreviewError("PowerPoint availability", "not installed"),
    ):
        with pytest.raises(DemoGenerationError) as first_error:
            generate_demo_project(deck, tmp_path / "workspace")
        with pytest.raises(DemoGenerationError) as second_error:
            generate_demo_project(deck, tmp_path / "workspace")

    assert first_error.value.project_id != second_error.value.project_id
    assert "PowerPoint availability" in first_error.value.reason
    assert (tmp_path / "workspace" / first_error.value.project_id).is_dir()
    assert (tmp_path / "workspace" / second_error.value.project_id).is_dir()
